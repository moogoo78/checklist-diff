"""Read a CSV or spreadsheet through a YAML field map.

Local checklists arrive as whatever columns their maintainer chose, often in
Chinese or another local language, and the columns change between releases. A
declarative map means onboarding a new source is a config file, not code:

    # map.yml
    fields:
      source_taxon_id: 學名編號
      scientific_name: 學名
      authorship: 命名者
      rank: 階層
      status: 分類狀態
      accepted_taxon_id: 接受名編號
      parent_taxon_id: 上階層編號
    classification:
      kingdom: 界
      family: 科名
    vernacular:
      - column: 中文名
        language: zh-TW
    status_values:          # optional, extends the built-in vocabulary
      接受: accepted
      同物異名: synonym

Only `scientific_name` is required. Every other field is optional, because
plenty of real checklists are a single column of names.
"""

from __future__ import annotations

import csv
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

import yaml

from checklistdiff.ingest.rows import SourceRow
from checklistdiff.models.enums import SourceFormat, TaxonomicStatus

log = logging.getLogger(__name__)

SIMPLE_FIELDS = (
    "source_taxon_id",
    "scientific_name",
    "authorship",
    "rank",
    "status",
    "accepted_taxon_id",
    "accepted_scientific_name",
    "parent_taxon_id",
    "parent_scientific_name",
)


class FieldMapError(ValueError):
    pass


@dataclass
class FieldMap:
    fields: dict[str, str]
    classification: dict[str, str] = field(default_factory=dict)
    vernacular: list[dict[str, str]] = field(default_factory=list)
    status_values: dict[str, str] = field(default_factory=dict)
    delimiter: str = ","
    encoding: str = "utf-8-sig"
    sheet: str | None = None

    @classmethod
    def from_yaml(cls, path: Path) -> "FieldMap":
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        fields = data.get("fields") or {}
        if "scientific_name" not in fields:
            raise FieldMapError(
                f"{path}: 'fields.scientific_name' is required — it is the only "
                "column ChecklistDiff cannot work without"
            )
        unknown = set(fields) - set(SIMPLE_FIELDS)
        if unknown:
            raise FieldMapError(
                f"{path}: unknown field(s) {sorted(unknown)}; "
                f"valid keys are {list(SIMPLE_FIELDS)}"
            )

        status_values = {
            str(k).strip().lower(): str(v).strip()
            for k, v in (data.get("status_values") or {}).items()
        }
        valid = {s.value for s in TaxonomicStatus}
        bad = {v for v in status_values.values() if v not in valid}
        if bad:
            raise FieldMapError(
                f"{path}: status_values maps to unknown status {sorted(bad)}; "
                f"valid targets are {sorted(valid)}"
            )

        return cls(
            fields=fields,
            classification=data.get("classification") or {},
            vernacular=data.get("vernacular") or [],
            status_values=status_values,
            delimiter=data.get("delimiter", ","),
            encoding=data.get("encoding", "utf-8-sig"),
            sheet=data.get("sheet"),
        )

    def validate_header(self, header: list[str], source: Path) -> None:
        """Fail fast on a column that the map names but the file lacks.

        Without this, a renamed column silently yields NULL for every row and the
        next diff reports the entire checklist as changed.
        """
        present = set(header)
        wanted: dict[str, str] = {}
        wanted.update(self.fields)
        wanted.update(self.classification)
        for v in self.vernacular:
            if col := v.get("column"):
                wanted[f"vernacular:{col}"] = col

        missing = {k: c for k, c in wanted.items() if c not in present}
        if missing:
            raise FieldMapError(
                f"{source}: mapped column(s) not in file: "
                + ", ".join(f"{k} -> {c!r}" for k, c in sorted(missing.items()))
                + f"\navailable columns: {sorted(present)}"
            )


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _iter_csv(path: Path, fmap: FieldMap) -> Iterator[dict[str, Any]]:
    with path.open("r", encoding=fmap.encoding, newline="") as fh:
        reader = csv.DictReader(fh, delimiter=fmap.delimiter)
        if reader.fieldnames:
            fmap.validate_header(list(reader.fieldnames), path)
        yield from reader


def _iter_xlsx(path: Path, fmap: FieldMap) -> Iterator[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[fmap.sheet] if fmap.sheet else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            return
        fmap.validate_header(header, path)
        for values in rows:
            if all(v is None for v in values):
                continue
            yield dict(zip(header, values, strict=False))
    finally:
        wb.close()


def detect_format(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return SourceFormat.XLSX.value
    return SourceFormat.CSV.value


def read_rows(path: Path, map_path: Path) -> Iterator[SourceRow]:
    """Yield `SourceRow`s from a CSV/XLSX file using the given field map."""
    fmap = FieldMap.from_yaml(map_path)
    fmt = detect_format(path)
    raw_rows = _iter_xlsx(path, fmap) if fmt == SourceFormat.XLSX.value else _iter_csv(path, fmap)

    for raw in raw_rows:
        get = lambda key: _clean(raw.get(fmap.fields[key])) if key in fmap.fields else None  # noqa: E731

        name = get("scientific_name")
        if not name:
            continue

        status_raw = get("status")
        if status_raw and (mapped := fmap.status_values.get(status_raw.lower())):
            status_raw = mapped

        classification = {
            rank: value
            for rank, column in fmap.classification.items()
            if (value := _clean(raw.get(column)))
        }

        vernacular = [
            {"name": value, "language": entry.get("language")}
            for entry in fmap.vernacular
            if (value := _clean(raw.get(entry.get("column", ""))))
        ]

        yield SourceRow(
            source_taxon_id=get("source_taxon_id"),
            scientific_name=name,
            authorship=get("authorship"),
            rank=get("rank"),
            status_raw=status_raw,
            accepted_taxon_id=get("accepted_taxon_id"),
            accepted_scientific_name=get("accepted_scientific_name"),
            parent_taxon_id=get("parent_taxon_id"),
            parent_scientific_name=get("parent_scientific_name"),
            classification=classification,
            vernacular=vernacular,
            raw={k: v for k, v in raw.items() if v not in (None, "")},
        )
